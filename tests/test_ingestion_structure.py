"""Stage 1 keeps what the document holds: DOCX tables, and PDF pages that are
scans inside an otherwise text PDF.

The PDFs are written by hand (no PDF library in the test environment): a text
page is a Helvetica content stream, a "scanned" page is one full-page
grayscale image and no text.
"""
import shutil

import docx
import pytest

from pipeline import stage1_ingestion as s1
from pipeline.stage1_ingestion import ingest

IOC = "185.220.101.45"


# ── DOCX ─────────────────────────────────────────────────────────────────────

def test_docx_tables_are_read_in_reading_order(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Indicators of compromise follow.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Type"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "C2 address"
    table.cell(1, 1).text = IOC
    doc.add_paragraph("End of report.")
    path = tmp_path / "r.docx"
    doc.save(str(path))

    lines = s1._read_docx(path).splitlines()

    assert lines == ["Indicators of compromise follow.", "Type | Value",
                     f"C2 address | {IOC}", "End of report."]
    assert IOC in ingest(str(path))


def test_docx_merged_and_nested_cells_are_read_once(tmp_path):
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "SUNBURST hashes"
    inner = table.cell(1, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "sha256"
    inner.cell(0, 1).text = "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855"
    path = tmp_path / "r.docx"
    doc.save(str(path))

    text = s1._read_docx(path)

    assert text.count("SUNBURST hashes") == 1
    assert "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855" in text


# ── PDF ──────────────────────────────────────────────────────────────────────

def _make_pdf(path, pages):
    """pages: ("text", str) or ("image", (width, height, 8-bit gray bytes))."""
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    catalog, pages_obj = add(b""), add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for kind, payload in pages:
        if kind == "text":
            ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
            for line in payload.split("\n"):
                esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                ops.append(f"({esc}) Tj T*")
            ops.append("ET")
            stream = "\n".join(ops).encode("latin-1")
            resources = f"<< /Font << /F1 {font} 0 R >> >>"
        else:
            w, h, data = payload
            img = add(b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace "
                      b"/DeviceGray /BitsPerComponent 8 /Length %d >>\nstream\n" % (w, h, len(data))
                      + data + b"\nendstream")
            stream = b"q 612 0 0 792 0 0 cm /Im1 Do Q"
            resources = f"<< /XObject << /Im1 {img} 0 R >> >>"
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(add(f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] "
                        f"/Resources {resources} /Contents {content} 0 R >>".encode()))
    objs[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
    objs[pages_obj - 1] = (f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] "
                           f"/Count {len(kids)} >>").encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root {catalog} 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))
    return path


TEXT_PAGE = ("text", "APT29 report, page one.\nThe actor deployed SUNBURST against targets\n"
                     "in several sectors during the campaign described below.")
BLANK_SCAN = ("image", (40, 52, bytes([255]) * (40 * 52)))


def test_a_scanned_page_inside_a_text_pdf_is_ocrd_in_place(tmp_path, monkeypatch):
    path = _make_pdf(tmp_path / "mixed.pdf", [TEXT_PAGE, BLANK_SCAN, TEXT_PAGE])
    calls = []
    monkeypatch.setattr(s1, "_ocr_page", lambda p, n: calls.append(n) or f"OCR page {n}: {IOC}")
    monkeypatch.setattr(s1, "_read_pdf_ocr", lambda p: pytest.fail("whole-document OCR"))

    text = s1._read_pdf(path)

    assert calls == [2]
    parts = text.split(f"OCR page 2: {IOC}")
    assert len(parts) == 2 and "page one" in parts[0] and "page one" in parts[1]


def test_a_sparse_page_without_an_image_is_not_ocrd(tmp_path, monkeypatch):
    path = _make_pdf(tmp_path / "cover.pdf", [("text", "Title"), TEXT_PAGE])
    monkeypatch.setattr(s1, "_ocr_page", lambda p, n: pytest.fail("OCR on a text page"))
    monkeypatch.setattr(s1, "_read_pdf_ocr", lambda p: pytest.fail("whole-document OCR"))

    assert "SUNBURST" in s1._read_pdf(path)


def test_a_fully_scanned_pdf_takes_the_whole_document_ocr(tmp_path, monkeypatch):
    path = _make_pdf(tmp_path / "scan.pdf", [BLANK_SCAN, BLANK_SCAN])
    monkeypatch.setattr(s1, "_read_pdf_ocr", lambda p: "whole-document OCR")
    assert s1._read_pdf(path) == "whole-document OCR"


def test_ocr_language_comes_from_the_environment(monkeypatch):
    monkeypatch.delenv("OCR_LANG", raising=False)
    assert s1.ocr_lang() == "eng"
    monkeypatch.setenv("OCR_LANG", "eng+fra")
    assert s1.ocr_lang() == "eng+fra"


@pytest.mark.skipif(not (s1._OCR_AVAILABLE and shutil.which("tesseract") and shutil.which("pdftoppm")),
                    reason="needs tesseract and poppler")
def test_real_ocr_recovers_an_ioc_from_a_scanned_page(tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", (1224, 1584), 255)
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default(size=56)
    except TypeError:        # Pillow < 10.1: no scalable default font
        pytest.skip("Pillow too old for a sized default font")
    draw.text((100, 300), f"C2 server {IOC}", fill=0, font=font)
    path = _make_pdf(tmp_path / "mixed.pdf",
                     [TEXT_PAGE, ("image", (img.width, img.height, img.tobytes()))])

    assert IOC in ingest(str(path))
