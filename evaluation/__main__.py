"""python -m evaluation <command> — the ADR-0060 protocol, step by step.

  prepare         registry of annotated layers, ATT&CK mapping, near-duplicates
  run             run the application's pipeline on a split (resumable)
  score           score a run: TTPs, entities, evidence
  ttp-errors      where a run's technique errors arise: never retrieved, not selected, origin of FPs
  compare         paired bootstrap between two runs on the same documents
  attack-diff     what changed between two ATT&CK releases, technique by technique
  release-effect  two runs' technique differences split into taxonomy and system effects
  retrieval       candidate recall per passage / document at several k (no LLM)
  support-sample  CSV of quotes for a person to judge
  preannotate     draft gold.json for an in-house report (to be corrected by hand)
  inhouse-run     run the pipeline on the in-house reports, from their original files
  inhouse-score   score an in-house run: IoCs, entities, techniques, relations, negatives
  hard-docs       list what makes candidate reports hard to ingest (tables, scans, columns)
  agreement       two independent annotations of one report, layer by layer

Outputs go to data/eval/ (git-ignored).  Tune on `dev`; run `test` once per
frozen configuration (see docs/eval/README.md).

The gold ids are mapped against one ATT&CK release (ADR-0087): a run's own by
default, `--attack-release` for another.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

from dotenv import load_dotenv

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from evaluation import annoctr, attack, dedup, evidence, inhouse, metrics  # noqa: E402
from pipeline import attack_data  # noqa: E402

DATA = _ROOT / "data"
OUT = DATA / "eval"
_LITERAL_ID = __import__("re").compile(r"\bT\d{4}(?:\.\d{3})?\b")
_TECH_ID = __import__("re").compile(r"^T\d{4}(?:\.\d{3})?$")


def _write(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def _docs(split: str, limit: int | None = None) -> list[annoctr.AnnoctrDoc]:
    docs = [d for d in annoctr.load(splits=(split,)) if "cyber" in d.layers]
    return docs[:limit] if limit else docs


# ── prepare ──────────────────────────────────────────────────────────────────

def cmd_prepare(args) -> None:
    docs = annoctr.load(include_general_only=True)
    reg = annoctr.registry(docs)
    _write(OUT / "annoctr-registry.json", reg)

    # The default track unless told otherwise, whatever ATTACK_RELEASE says.
    cat = attack.Catalogue.for_release(args.attack_release or "", data_dir=DATA)
    gold = {t for d in docs for t in d.technique_ids}
    mapping = cat.mapping(gold)
    version = cat.summary()["release"]
    _write(OUT / ("attack-mapping.json" if version == attack.default_release()
                  else f"attack-mapping-{version}.json"),
           {"catalogue": cat.summary(), "rows": [r.as_dict() for r in mapping.values()]})

    groups = {d.doc_id: (d.split, d.text) for d in docs}
    for extra in args.corpus or []:
        for p in sorted(Path(extra).glob("*.txt")):
            groups[f"{Path(extra).name}/{p.stem}"] = (Path(extra).name, p.read_text(encoding="utf-8"))
    dups = dedup.near_duplicates(groups, threshold=args.threshold)
    _write(OUT / "near-duplicates.json", dups)

    by_split: dict[str, int] = {}
    for row in reg:
        by_split[row["split"]] = by_split.get(row["split"], 0) + 1
    statuses: dict[str, int] = {}
    for res in mapping.values():
        statuses[res.status] = statuses.get(res.status, 0) + 1
    print(f"documents: {by_split}")
    print(f"gold technique ids: {len(gold)} — {statuses}; "
          f"not scorable: {sum(1 for res in mapping.values() if not res.in_pipeline)}")
    for res in mapping.values():
        if res.status != attack.ACTIVE:
            print(f"  {res.gold_id:<11} {res.status:<11} -> {res.canonical}")
    cross = [d for d in dups if d["cross"]]
    print(f"near-duplicate pairs (>= {args.threshold}): {len(dups)}, across groups: {len(cross)}")
    for d in cross[:10]:
        print(f"  {d['a']} [{d['group_a']}] ~ {d['b']} [{d['group_b']}] "
              f"jaccard={d['jaccard']} containment={d['containment']}")
    print(f"written to {OUT}")


# ── run ──────────────────────────────────────────────────────────────────────

def _predictions(result) -> dict:
    """What the application would show for this document."""
    from models.schemas import EntityType

    llm = result.llm_result
    ttp_ids = {(e.mitre_id or e.value).upper() for e in result.entities
               if e.entity_type == EntityType.TTP}
    ttps = []
    if llm is not None:
        for t in llm.ttps:
            if t.mitre_id:
                ttp_ids.add(t.mitre_id.upper())
                ttps.append({"mitre_id": t.mitre_id.upper(), "name": t.technique_name,
                             "evidence_text": t.evidence_text,
                             "evidence_label": getattr(t.evidence_label, "value", t.evidence_label)})
    entities = [{"type": e.entity_type.value, "value": e.value, "source": e.source,
                 "confidence": e.confidence, "mitre_id": e.mitre_id} for e in result.entities]
    if llm is not None:
        for etype, names in (("malware", llm.malware_families), ("threat_actor", llm.threat_actors),
                             ("tool", llm.tools), ("location", llm.targeted_countries),
                             ("identity", llm.targeted_sectors)):
            entities += [{"type": etype, "value": n, "source": "llm", "confidence": None,
                          "mitre_id": None} for n in names]
    relations = []
    if llm is not None:
        relations = [{"source": r.source_value, "type": r.relationship_type, "target": r.target_value,
                      "evidence_text": r.evidence_text,
                      "evidence_label": getattr(r.evidence_label, "value", r.evidence_label)}
                     for r in llm.relationships]
    # ADR-0072 select mode: what 3f left undecided, and what it chose from.  The
    # quote of a refused selection is kept: without it a refusal cannot be told
    # apart as a missing, a too-short or a non-verbatim quote.
    review = ([{"mitre_id": rv.attack_id, "reason": rv.reason, "sources": rv.sources,
                "evidence_quote": rv.evidence_quote}
               for rv in llm.ttp_review] if llm is not None else [])
    candidates = sorted({c.attack_id for cs in getattr(result, "ttp_candidates", []) for c in cs})
    # ADR-0082: the claims 3d held back.  Not in `relations` — they do not
    # ship unless an analyst accepts them — but countable, with the reason.
    rel_review = ([{"source": h.source_value, "type": h.relationship_type, "target": h.target_value,
                    "evidence_text": h.evidence_text, "reason": h.reason}
                   for h in llm.rel_review] if llm is not None else [])
    return {"ttp_ids": sorted(i for i in ttp_ids if _TECH_ID.match(i)), "ttps": ttps,
            "ttp_review": review, "ttp_candidates": candidates,
            "entities": entities, "relations": relations, "rel_review": rel_review,
            "bundle_relations": _bundle_relations(result.bundle)}


def _bundle_relations(bundle) -> list[dict] | None:
    """The relationships the bundle ships, endpoints resolved to names — what
    an analyst receives, 4b / 4c edges included.  None without a bundle."""
    if bundle is None:
        return None
    objs = json.loads(bundle.serialize())["objects"]

    def name(o):
        return o.get("name") or o.get("value") or o.get("pattern") or o.get("id")

    by_id = {o["id"]: name(o) for o in objs}
    return [{"source": by_id.get(o["source_ref"], o["source_ref"]), "type": o["relationship_type"],
             "target": by_id.get(o["target_ref"], o["target_ref"]),
             "evidence_label": o.get("x_evidence_label"), "policy_rule": o.get("x_policy_rule")}
            for o in objs if o.get("type") == "relationship"]


def _stay_on_the_release(meta_path: Path, name: str) -> None:
    """The release this process runs on, checked before any document: an
    unknown ATTACK_RELEASE stops the command here instead of marking every
    document an error, and a run that already began resumes on the release its
    run.json recorded, not on another track than the documents already done
    (ADR-0087)."""
    version = attack_data.current().version
    if meta_path.exists():
        began = run_release(json.loads(meta_path.read_text(encoding="utf-8")))
        if began != version:
            sys.exit(f"{name} began on ATT&CK {began} and ATTACK_RELEASE selects {version}: "
                     f"resume it with ATTACK_RELEASE={began}, or start another run")


def cmd_run(args) -> None:
    from api.run_config import build_manifest
    from pipeline.orchestrator import Document, RunOptions, StageRequired, run_document

    run_dir = OUT / "runs" / args.name
    meta_path = run_dir / "run.json"
    _stay_on_the_release(meta_path, args.name)
    disabled = {"1f", "4", "5"} | {s.strip() for s in (args.disable or "").split(",") if s.strip()}
    # The procedures MITRE wrote from AnnoCTR reports stay out of the select
    # path's retrieval corpus during an evaluation (ADR-0072).
    if not os.environ.get("TTP_RETRIEVAL_EXCLUDE_CITED"):
        excl = OUT / "annoctr-cited-exclude.txt"
        excl.parent.mkdir(parents=True, exist_ok=True)
        excl.write_text("\n".join(annoctr_cited_exclusions()) + "\n", encoding="utf-8")
        os.environ["TTP_RETRIEVAL_EXCLUDE_CITED"] = str(excl)
    options = RunOptions.from_env(disabled=disabled,
                                  required=set() if args.allow_degraded else {"3"})
    docs = _docs(args.split, args.limit)
    if not meta_path.exists():
        _write(meta_path, {"name": args.name, "split": args.split, "started": time.strftime("%Y-%m-%dT%H:%M:%S"),
                           "disabled": sorted(options.disabled), "required": sorted(options.required),
                           "consensus": options.consensus, "document_relations": options.document_relations,
                           "ttp_mode": options.ttp_mode,
                           "manifest": build_manifest()})
    consecutive_required = 0
    for i, d in enumerate(docs, 1):
        out = run_dir / "docs" / f"{d.doc_id}.json"
        if out.exists():
            continue
        t0 = time.monotonic()
        try:
            result = run_document(Document(text=d.text, original_filename=f"{d.doc_id}.txt"), options)
            record = {"doc_id": d.doc_id, "seconds": round(time.monotonic() - t0, 1),
                      "stages": result.stage_report(), **_predictions(result)}
            consecutive_required = 0
        except StageRequired as exc:
            # One report whose LLM answers were all unusable is recorded and
            # left out of the score; two in a row is the setup (a model name
            # the server does not serve, a server down): stop.
            record = {"doc_id": d.doc_id, "seconds": round(time.monotonic() - t0, 1),
                      "error": f"StageRequired: {exc}"}
            consecutive_required += 1
            if consecutive_required >= 2:
                _write(out, record)
                print(f"stopped: {exc}", file=sys.stderr)
                sys.exit(2)
        except attack_data.UnknownAttackRelease:
            raise   # the setting is wrong for the whole run, not for this document
        except Exception as exc:   # one bad document must not lose the run
            record = {"doc_id": d.doc_id, "seconds": round(time.monotonic() - t0, 1),
                      "error": f"{type(exc).__name__}: {exc}"}
        _write(out, record)
        print(f"[{i}/{len(docs)}] {d.doc_id} {record['seconds']}s "
              f"{'ERROR ' + record['error'] if 'error' in record else str(len(record['ttp_ids'])) + ' techniques'}",
              flush=True)


# ── score ────────────────────────────────────────────────────────────────────

def _read_meta(name: str) -> dict:
    return json.loads((OUT / "runs" / name / "run.json").read_text(encoding="utf-8"))


def _load_run(name: str) -> tuple[dict, dict[str, dict]]:
    run_dir = OUT / "runs" / name
    meta = _read_meta(name)
    preds = {p.stem: json.loads(p.read_text(encoding="utf-8")) for p in (run_dir / "docs").glob("*.json")}
    return meta, preds


def run_release(meta: dict) -> str:
    """The ATT&CK release a run's pipeline ran on.  A run from before the
    manifest recorded it (no `attack` block) ran on the default track."""
    recorded = ((meta.get("manifest") or {}).get("attack") or {}).get("release")
    return str(recorded) if recorded else attack.default_release()


def _modified_files(name: str, meta: dict) -> list[str]:
    """The ATT&CK files that no longer matched their track's manifest when the
    run began (`attack.modified_files` of its manifest).  They are said on
    stderr: the release the run announces is then not proven, but the score
    stands."""
    files = list(((meta.get("manifest") or {}).get("attack") or {}).get("modified_files") or [])
    if files:
        print(f"WARNING: {name} ran on ATT&CK {run_release(meta)} files that differ from its manifest "
              f"({', '.join(files)}): the release it announces is not proven", file=sys.stderr)
    return files


def _shared_release(a: str, b: str, option: str | None) -> str:
    """The release two runs are scored against: the one asked for, or the one
    both ran on.  The gold has to be mapped the same way on both sides."""
    if option:
        return option
    ra, rb = run_release(_read_meta(a)), run_release(_read_meta(b))
    if ra != rb:
        sys.exit(f"{a} ran on ATT&CK {ra} and {b} on {rb}: pass --attack-release to say which "
                 "release the gold is mapped against, for both")
    return ra


def _canonical_fn(release: str | None = None):
    cat = attack.Catalogue.for_release(release, data_dir=DATA)
    cache: dict[str, str | None] = {}

    def canonical(tid: str) -> str | None:
        if tid not in cache:
            r = cat.resolve(tid)
            cache[tid] = r.canonical if r.in_pipeline else None
        return cache[tid]
    return canonical, cat


def per_doc_scores(name: str, release: str | None = None) -> dict:
    """`release`: the ATT&CK release to map the gold with; the run's own when None."""
    meta, preds = _load_run(name)
    canonical, cat = _canonical_fn(release or run_release(meta))
    docs = {d.doc_id: d for d in _docs(meta["split"])}
    rows = {}
    for doc_id, p in sorted(preds.items()):
        d = docs.get(doc_id)
        if d is None or "error" in p:
            continue
        ttp = metrics.score_ttps(set(p["ttp_ids"]), d.technique_kinds(), canonical)
        # Sensitivity check: AnnoCTR leaves most ids of a vendor's own ATT&CK
        # table unannotated (15 of 23 on dev), so an id the report writes out
        # counts as a false positive.  Same score with those ids added to gold.
        literal = {m.upper(): {"explicit"} for m in _LITERAL_ID.findall(d.text)}
        ttp_literal = metrics.score_ttps(set(p["ttp_ids"]), {**literal, **d.technique_kinds()}, canonical)
        gold_ents = [(e.etype, e.label, e.surfaces) for e in d.entities.values()]
        types = set(d.evaluable()["entity_types"])
        pred_ents = [(e["type"], e["value"]) for e in p["entities"] if e["type"] in types]
        rows[doc_id] = {
            "ttp": ttp,
            "ttp_with_literal_ids": ttp_literal,
            "entities_strict": metrics.score_entities(pred_ents, gold_ents),
            "entities_lenient": metrics.score_entities(pred_ents, gold_ents, lenient=True),
        }
    return {"meta": meta, "rows": rows, "preds": preds, "docs": docs, "canonical": canonical,
            "catalogue": cat}


def cmd_score(args) -> None:
    s = per_doc_scores(args.name, args.attack_release)
    modified = _modified_files(args.name, s["meta"])
    rows, preds, docs, canonical = s["rows"], s["preds"], s["docs"], s["canonical"]
    tech = [r["ttp"].technique for r in rows.values()]
    sub = [r["ttp"].subtechnique for r in rows.values()]
    exp = [r["ttp"].explicit_recall for r in rows.values()]
    imp = [r["ttp"].implicit_recall for r in rows.values()]

    ent_types = sorted({t for r in rows.values() for t in r["entities_strict"]})
    ents = {t: {"strict": metrics.micro(r["entities_strict"].get(t, metrics.Counts())
                                        for r in rows.values()).as_dict(),
                "lenient": metrics.micro(r["entities_lenient"].get(t, metrics.Counts())
                                         for r in rows.values()).as_dict()}
            for t in ent_types}

    ev_rows = []
    for doc_id in rows:
        d = docs[doc_id]
        spans: dict[str, list[tuple[int, int]]] = {}
        for m in d.techniques:
            c = canonical(m.label)
            if c and m.start is not None and m.end is not None:
                spans.setdefault(c.split(".", 1)[0], []).append((m.start, m.end))
        ev_rows += evidence.check(doc_id, d.text, preds[doc_id]["ttps"], spans, canonical)

    stages: dict[str, dict[str, int]] = {}
    llm_health: dict[str, int] = {}
    for p in preds.values():
        for o in p.get("stages", []):
            stages.setdefault(o["stage"], {}).setdefault(o["status"], 0)
            stages[o["stage"]][o["status"]] += 1
            if o["stage"] in ("3", "3d", "3f"):
                for k, v in o.get("counts", {}).items():
                    if k.startswith(("extraction_", "provider_", "ok", "unparsed", "failed", "truncated")):
                        key = k if o["stage"] == "3" else f"{o['stage']}_{k}"
                        llm_health[key] = llm_health.get(key, 0) + v

    def ratio(pairs):
        found, gold = sum(a for a, _ in pairs), sum(b for _, b in pairs)
        return {"found": found, "gold": gold, "recall": round(found / gold, 4) if gold else None}

    report = {
        "run": args.name, "split": s["meta"]["split"], "documents_scored": len(rows),
        "documents_failed": sorted(k for k, p in preds.items() if "error" in p),
        "config": {k: s["meta"][k] for k in ("disabled", "consensus", "document_relations")},
        "llm": s["meta"]["manifest"].get("llm"),
        "attack_mapping": s["catalogue"].summary(),
        **({"attack_modified_files": modified} if modified else {}),
        "ttp": {
            "technique_micro": metrics.micro(tech).as_dict(),
            "technique_per_document": metrics.macro(tech),
            "subtechnique_micro": metrics.micro(sub).as_dict(),
            "explicit_recall": ratio(exp), "implicit_only_recall": ratio(imp),
            "gold_dropped_by_mapping": sorted({g for r in rows.values() for g in r["ttp"].dropped_gold}),
            "technique_micro_gold_plus_literal_ids": metrics.micro(
                r["ttp_with_literal_ids"].technique for r in rows.values()).as_dict(),
        },
        "entities": ents,
        "evidence": evidence.summarise(ev_rows),
        "stages": stages,
        "llm_health": llm_health,
        "seconds_per_document": round(sum(p.get("seconds", 0) for p in preds.values()) / max(1, len(preds)), 1),
    }
    release, own = report["attack_mapping"]["release"], run_release(s["meta"])
    _write(OUT / "runs" / args.name / ("score.json" if release == own else f"score-attack-{release}.json"),
           report)
    t = report["ttp"]
    print(f"{args.name} [{report['split']}] {len(rows)} documents "
          f"({len(report['documents_failed'])} failed), {report['seconds_per_document']} s/doc")
    if release != own:
        print(f"  gold mapped with ATT&CK {release}; the run used {own}")
    print(f"  TTP technique  micro P/R/F1 {t['technique_micro']['precision']:.3f} "
          f"{t['technique_micro']['recall']:.3f} {t['technique_micro']['f1']:.3f}   "
          f"per-document F1 {t['technique_per_document']['f1']:.3f}")
    print(f"  TTP recall     explicit {t['explicit_recall']['recall']}  implicit-only "
          f"{t['implicit_only_recall']['recall']}")
    lit = t["technique_micro_gold_plus_literal_ids"]
    print(f"  TTP technique  gold + ids written in the report: P/R/F1 {lit['precision']:.3f} "
          f"{lit['recall']:.3f} {lit['f1']:.3f}  (sensitivity check, not the primary score)")
    for et, v in ents.items():
        print(f"  {et:<13} strict P/R/F1 {v['strict']['precision']:.3f} {v['strict']['recall']:.3f} "
              f"{v['strict']['f1']:.3f}   lenient F1 {v['lenient']['f1']:.3f}")
    print(f"  evidence       {report['evidence']}")
    skipped = {k: v for k, v in stages.items() if set(v) - {"ran"}}
    if skipped:
        print(f"  stages not run: {skipped}")
    print(f"  LLM health     {llm_health}")


# ── ttp-errors ───────────────────────────────────────────────────────────────

def cmd_ttp_errors(args) -> None:
    """Split a run's technique errors by where they arise (ADR-0072): never
    retrieved, retrieved but not selected, and false positives by origin."""
    meta, preds = _load_run(args.name)
    modified = _modified_files(args.name, meta)
    canonical, catalogue = _canonical_fn(args.attack_release or run_release(meta))
    docs = {d.doc_id: d for d in _docs(meta["split"])}
    techniques: dict[str, dict[str, int]] = {k: {} for k in metrics.TTP_ERROR_CATEGORIES}
    reasons: dict[str, int] = {}
    per_doc = {}
    for doc_id, p in sorted(preds.items()):
        d = docs.get(doc_id)
        if d is None or "error" in p:
            continue
        review = {rv["mitre_id"]: rv.get("reason") or "" for rv in p.get("ttp_review") or []}
        split, why = metrics.ttp_error_split(p["ttp_ids"], d.technique_kinds(),
                                             p.get("ttp_candidates") or [], review, canonical)
        per_doc[doc_id] = split
        for cat, ids in split.items():
            for tid in ids:
                techniques[cat][tid] = techniques[cat].get(tid, 0) + 1
        for k, v in why.items():
            reasons[k] = reasons.get(k, 0) + v
    totals = {cat: sum(ids.values()) for cat, ids in techniques.items()}
    gold_n = totals["found"] + totals["retrieved_not_selected"] + totals["never_retrieved"]
    report = {
        "run": args.name, "documents": len(per_doc), "totals": totals, "gold": gold_n,
        "attack_mapping": catalogue.summary(),
        **({"attack_modified_files": modified} if modified else {}),
        "candidate_ceiling": round((totals["found"] + totals["retrieved_not_selected"]) / gold_n, 4)
        if gold_n else None,
        "review_reasons_of_missed_gold": reasons,
        "most_frequent": {cat: sorted(ids.items(), key=lambda kv: (-kv[1], kv[0]))[:10]
                          for cat, ids in techniques.items() if cat != "found"},
        "per_document": per_doc,
    }
    release = report["attack_mapping"]["release"]
    _write(OUT / "runs" / args.name / ("ttp-errors.json" if release == run_release(meta)
                                       else f"ttp-errors-attack-{release}.json"), report)
    print(f"{args.name}: {len(per_doc)} documents, {gold_n} gold parent techniques")
    for cat in metrics.TTP_ERROR_CATEGORIES:
        print(f"  {cat:<24} {totals[cat]}")
    print(f"  candidate ceiling        {report['candidate_ceiling']}")
    if reasons:
        print(f"  held back by review      {reasons}")


# ── compare ──────────────────────────────────────────────────────────────────

def cmd_compare(args) -> None:
    release = _shared_release(args.a, args.b, args.attack_release)
    a, b = per_doc_scores(args.a, release), per_doc_scores(args.b, release)
    modified = {"a": _modified_files(args.a, a["meta"]), "b": _modified_files(args.b, b["meta"])}
    common = sorted(set(a["rows"]) & set(b["rows"]))
    if not common:
        sys.exit("no document scored in both runs")

    def series(s, doc_ids):
        if args.task == "ttp":
            return [s["rows"][d]["ttp"].technique for d in doc_ids]
        return [s["rows"][d]["entities_strict"].get(args.task, metrics.Counts()) for d in doc_ids]

    out = {"a": args.a, "b": args.b, "task": args.task, "documents": len(common),
           "attack_mapping": a["catalogue"].summary(),
           **({"attack_modified_files": modified} if modified["a"] or modified["b"] else {}), "results": {}}
    for m in ("f1", "precision", "recall"):
        out["results"][m] = metrics.paired_bootstrap(series(a, common), series(b, common), metric=m,
                                                     n=args.resamples)
    suffix = f"__attack-{out['attack_mapping']['release']}" if args.attack_release else ""
    _write(OUT / "compare" / f"{args.a}__vs__{args.b}__{args.task}{suffix}.json", out)
    print(f"{args.b} − {args.a} on {len(common)} documents ({args.task}):")
    for m, r in out["results"].items():
        print(f"  {m:<9} {r['observed_diff']:+.4f}  95% CI [{r['ci95'][0]:+.4f}, {r['ci95'][1]:+.4f}]  "
              f"P(not better) {r['p_not_better']}")


# ── attack-diff, release-effect ──────────────────────────────────────────────

def cmd_attack_diff(args) -> None:
    old, new = attack.release_version(args.old), attack.release_version(args.new)
    diff = attack.release_diff(attack.release_dir(old, DATA), attack.release_dir(new, DATA))
    path = OUT / f"attack-diff-{old}-{new}.json"
    _write(path, diff)
    print(f"ATT&CK {old} -> {new}")
    for domain, r in diff.items():
        print(f"  {domain:<18} active {r['active']['old']} -> {r['active']['new']}, "
              f"new {len(r['new_active'])}, no longer active {len(r['no_longer_active'])}, "
              f"renamed {len(r['renamed'])}, description changed {len(r['description_changed'])}, "
              f"procedures {r['procedures']['old']} -> {r['procedures']['new']}")
        if r["new_active"]:
            print(f"    new: {', '.join(r['new_active'])}")
        if r["no_longer_active"]:
            print(f"    no longer active: {', '.join(r['no_longer_active'])}")
    print(f"written to {path}")


_NAMED_TYPES = ("malware", "threat_actor", "tool")


def _ttp_parents(ids, canonical) -> set[str]:
    """The parent-level ids of one document in the scoring taxonomy, as
    `metrics.score_ttps` builds them."""
    return {c.split(".", 1)[0] for i in ids if (c := canonical(i)) is not None}


def _gazetteer_entities(pred: dict, by_source: bool) -> set[tuple]:
    """What the gazetteer gave one document: the entities it sourced or, for
    runs that record no entity source, every malware, threat actor and tool."""
    return {(e["type"], e["value"].lower(), e.get("mitre_id")) for e in pred["entities"]
            if (e.get("source") == "gazetteer" if by_source else e["type"] in _NAMED_TYPES)}


def _only(a: set, b: set) -> dict | None:
    """What one side has and the other has not; None when they agree."""
    return None if a == b else {"only_a": sorted(a - b, key=str), "only_b": sorted(b - a, key=str)}


def _compared_releases(a: str, b: str, old: str | None, new: str | None) -> tuple[str, str]:
    """The two releases `release-effect` compares: those the runs ran on.  Two
    runs on the same release say nothing about which two to compare, so `--old`
    and `--new` are required then; otherwise an option has to be the release of
    the run it stands for."""
    ra, rb = run_release(_read_meta(a)), run_release(_read_meta(b))
    if ra == rb:
        if not (old and new):
            sys.exit(f"{a} and {b} both ran on ATT&CK {ra}: pass --old and --new to say which two "
                     "releases to compare")
        return attack.release_version(old), attack.release_version(new)
    for option, value, expected, run in (("--old", old, ra, a), ("--new", new, rb, b)):
        if value and attack.release_version(value) != expected:
            sys.exit(f"{option} {value} but {run} ran on ATT&CK {expected}")
    return ra, rb


def cmd_release_effect(args) -> None:
    """Two runs on two ATT&CK releases: how much of their parent-level technique
    differences the taxonomy explains (ADR-0087).  Each id of Pa △ Pb is classed
    by what changed between the releases (`attack.release_change`); per class,
    the change in TP / FP / FN, and the F1 A would have if only that class came
    from B.  Documents whose gazetteer entities or candidates differ are listed:
    the rest of the difference is not the taxonomy's."""
    old, new = _compared_releases(args.a, args.b, args.old, args.new)
    release = _shared_release(args.a, args.b, args.attack_release)
    a, b = per_doc_scores(args.a, release), per_doc_scores(args.b, release)
    modified = {"a": _modified_files(args.a, a["meta"]), "b": _modified_files(args.b, b["meta"])}
    common = sorted(set(a["rows"]) & set(b["rows"]))
    if not common:
        sys.exit("no document scored in both runs")
    classify = attack.release_change(
        attack.release_diff(attack.release_dir(old, DATA), attack.release_dir(new, DATA)))
    by_source = all("source" in e for s in (a, b) for d in common for e in s["preds"][d]["entities"])

    counts_a, counts_b = [], []
    classes: dict[str, dict] = {c: {"ids": set(), "differences": 0, "hybrid": []} for c in attack.CHANGE_CLASSES}
    detail = []
    for doc_id in common:
        gold = _ttp_parents(a["docs"][doc_id].technique_kinds(), a["canonical"])
        pa, pb = (_ttp_parents(s["preds"][doc_id]["ttp_ids"], s["canonical"]) for s in (a, b))
        for s, ids in ((a, pa), (b, pb)):
            if metrics.set_counts(ids, gold) != s["rows"][doc_id]["ttp"].technique:
                raise RuntimeError(f"{doc_id}: parent-level sets differ from metrics.score_ttps")
        counts_a.append(metrics.set_counts(pa, gold))
        counts_b.append(metrics.set_counts(pb, gold))
        why = {t: classify(t) for t in pa | pb}
        for t in pa ^ pb:
            classes[why[t]]["ids"].add(t)
            classes[why[t]]["differences"] += 1
        for c, row in classes.items():
            row["hybrid"].append(metrics.set_counts(
                {t for t in pa if why[t] != c} | {t for t in pb if why[t] == c}, gold))
        gazetteer = _only(*(_gazetteer_entities(s["preds"][doc_id], by_source) for s in (a, b)))
        candidates = _only(*(set(s["preds"][doc_id].get("ttp_candidates") or []) for s in (a, b)))
        if pa != pb or gazetteer or candidates:
            detail.append({"doc_id": doc_id,
                           "ttp_differences": [{"id": t, "class": why[t], "in": "a" if t in pa else "b",
                                                "gold": t in gold} for t in sorted(pa ^ pb)],
                           "gazetteer_entities": gazetteer, "candidates": candidates})

    total_a, total_b = metrics.micro(counts_a), metrics.micro(counts_b)
    per_class = {}
    for c, row in classes.items():
        h = metrics.micro(row["hybrid"])
        per_class[c] = {"n_ids": len(row["ids"]), "n_differences": row["differences"],
                        "delta_tp": h.tp - total_a.tp, "delta_fp": h.fp - total_a.fp,
                        "delta_fn": h.fn - total_a.fn, "f1_hybrid": round(h.f1, 4),
                        "ids": sorted(row["ids"])}
    differ = [d for d in detail if d["ttp_differences"]]
    report = {
        "a": args.a, "b": args.b, "releases": {"a": run_release(a["meta"]), "b": run_release(b["meta"]),
                                               "old": old, "new": new},
        "attack_mapping": a["catalogue"].summary(), "level": "parent technique",
        **({"attack_modified_files": modified} if modified["a"] or modified["b"] else {}),
        "documents": len(common),
        "f1_micro": {"a": round(total_a.f1, 4), "b": round(total_b.f1, 4)},
        "delta_total": {"tp": total_b.tp - total_a.tp, "fp": total_b.fp - total_a.fp,
                        "fn": total_b.fn - total_a.fn},
        "classes": per_class,
        "entities_compared": "entities whose source is the gazetteer" if by_source else
        "malware, threat_actor and tool entities: the runs record no entity source",
        "documents_with_ttp_difference": len(differ),
        "documents_gazetteer_differs": sum(1 for d in detail if d["gazetteer_entities"]),
        "documents_candidates_differ": sum(1 for d in detail if d["candidates"]),
        "documents_ttp_difference_and_upstream_difference": sum(
            1 for d in differ if d["gazetteer_entities"] or d["candidates"]),
        "documents_detail": detail,
    }
    path = OUT / "compare" / f"{args.a}__vs__{args.b}__release-effect.json"
    _write(path, report)
    print(f"{args.b} − {args.a} on {len(common)} documents, parent techniques, gold mapped with ATT&CK "
          f"{report['attack_mapping']['release']} (releases {old} -> {new})")
    print(f"  micro F1  a {report['f1_micro']['a']:.4f}  b {report['f1_micro']['b']:.4f}   "
          f"TP {report['delta_total']['tp']:+d}  FP {report['delta_total']['fp']:+d}  "
          f"FN {report['delta_total']['fn']:+d}")
    print(f"  {'class':<18} {'ids':>4} {'diffs':>6} {'dTP':>5} {'dFP':>5} {'dFN':>5}  F1 if only this from b")
    for c, r in per_class.items():
        print(f"  {c:<18} {r['n_ids']:>4} {r['n_differences']:>6} {r['delta_tp']:>+5d} {r['delta_fp']:>+5d} "
              f"{r['delta_fn']:>+5d}  {r['f1_hybrid']:.4f}")
    print(f"  documents: technique differs {report['documents_with_ttp_difference']}, gazetteer entities differ "
          f"{report['documents_gazetteer_differs']}, candidates differ {report['documents_candidates_differ']}; "
          f"technique and upstream both differ {report['documents_ttp_difference_and_upstream_difference']}")
    print(f"  compared entities: {report['entities_compared']}")
    print(f"written to {path}")


# ── retrieval ────────────────────────────────────────────────────────────────

def annoctr_cited_exclusions() -> list[str]:
    """URL fragments of every AnnoCTR report (its slug), all splits.

    MITRE cites some of these blog posts in its procedure examples — 48 in the
    2026-10 bundle, 28 of them from test reports.  A retrieval corpus that
    keeps them would hand the retriever sentences written from the report it
    is scored on (ADR-0072)."""
    out = []
    for split in annoctr.SPLITS:
        folder = annoctr.DEFAULT_ROOT / "text" / split
        for p in sorted(folder.glob("*.txt")) if folder.is_dir() else []:
            parts = p.stem.split("_", 2)
            if len(parts) == 3 and len(parts[2]) >= 12:
                out.append(parts[2])
    return sorted(set(out))


def _chunk_recall(docs, retriever, gate: bool, spec: str, canonical, cand_id) -> dict:
    """Recall of the candidate lists the selector gets, one per Stage 1 chunk
    (`spec` = "k_per_passage:cap")."""
    from evaluation import retrieval as rv
    from pipeline.orchestrator import chunk_size_for
    from pipeline.stage1_ingestion import chunk_text
    from pipeline.ttp_retrieval import chunk_candidates_with, split_passages

    k_s, _, cap_s = spec.partition(":")
    k, cap = int(k_s), int(cap_s or 0)
    out = []
    for d in docs:
        chunks = chunk_text(d.text, max_chars=chunk_size_for(len(d.text)))
        cands = chunk_candidates_with(retriever, chunks, k_per_passage=k, max_per_chunk=cap or None,
                                      keyword_gate=gate)
        passages = split_passages(d.text)
        gold: dict[str, list[str]] = {}
        gold_ids: set[str] = set()
        for m in d.techniques:
            c = canonical(m.label)
            if not c:
                continue
            gold_ids.add(c)
            if m.start is not None and m.end is not None:
                i = rv.passage_of([(p.start, p.end) for p in passages], m.start, m.end)
                if i is not None:
                    gold.setdefault(c, []).append(passages[i].text)
        out.append(rv.DocChunks(gold, gold_ids, [" ".join(ch.split()) for ch in chunks],
                                [[cand_id(x.attack_id) for x in cs] for cs in cands]))
    return {"k_per_passage": k, "cap": cap, **rv.chunk_row(out)}


def cmd_retrieval(args) -> None:
    from evaluation import retrieval as rv
    from pipeline import stage2c_ttp_semantic as s2c
    from pipeline import ttp_retrieval as tr

    canonical, cat = _canonical_fn(args.attack_release)

    def cand_id(t: str) -> str:
        # Candidates may come from a newer bundle than the pipeline index:
        # follow revocations, never drop a candidate for not being indexed.
        return cat.resolve(t).canonical or t.upper()

    ks = sorted({int(k) for k in args.k.split(",")})
    corpora = [c.strip() for c in args.corpus.split(",")]
    methods = [m.strip() for m in args.method.split(",")]
    gates = [g.strip() == "on" for g in args.gate.split(",")]
    excl = [] if args.keep_cited else annoctr_cited_exclusions()
    docs = _docs(args.split, args.limit)
    path = Path(args.corpus_dir) / "attack_retrieval_corpus.json" if args.corpus_dir else None
    emb = Path(args.corpus_dir) / "attack_retrieval_embeddings.npy" if args.corpus_dir else None
    model = None
    results = []
    for corpus_kind in corpora:
        corpus = tr.load_corpus(corpus_kind, path=path, emb_path=emb, exclude_cited=excl)
        if corpus is None:
            print(f"corpus {corpus_kind}: unavailable (python scripts/build_indexes.py --only retrieval)")
            continue
        for method in methods:
            if method != "bm25" and model is None:
                model = s2c._load_model()
                if model is None:
                    sys.exit("dense retrieval needs the embedding model (SKIP_HEAVY_MODELS?)")
            retriever = tr.Retriever(corpus, method, model if method != "bm25" else None)
            for gate in gates:
                rankings = []
                for d in docs:
                    t0 = time.monotonic()
                    passages = tr.split_passages(d.text)
                    kept = tr.gate_passages(passages, keyword_gate=gate)
                    kept_set = set(kept)
                    kept_idx = {i for i, p in enumerate(passages) if p in kept_set}
                    ranked_lists = retriever.rank(kept, k=max(ks))
                    by_pos = {(p.start, p.end): [cand_id(c.attack_id) for c in row]
                              for p, row in zip(kept, ranked_lists, strict=True)}
                    gold: dict[str, list[tuple[int, int]]] = {}
                    unlocated: set[str] = set()
                    for m in d.techniques:
                        c = canonical(m.label)
                        if not c:
                            continue
                        if m.start is None or m.end is None:
                            unlocated.add(c)
                        else:
                            gold.setdefault(c, []).append((m.start, m.end))
                    rankings.append(rv.DocRanking(
                        d.doc_id, len(passages), [(p.start, p.end) for p in passages], kept_idx,
                        {i: by_pos[(p.start, p.end)] for i, p in enumerate(passages) if i in kept_idx},
                        gold, unlocated - set(gold), time.monotonic() - t0))
                rows = [rv.recall_row(rankings, k) for k in ks]
                cfg = {"corpus": corpus_kind, "corpus_version": corpus.version,
                       "entries": len(corpus.ids), "method": method,
                       "keyword_gate": "on" if gate else "off",
                       "coverage": rv.coverage(rankings), "table": rows,
                       "buckets": {str(k): rv.buckets(rankings, k) for k in (min(ks), max(ks))},
                       "chunks": [_chunk_recall(docs, retriever, gate, kc, canonical, cand_id)
                                  for kc in (args.chunk or "").split(",") if kc.strip()]}
                results.append(cfg)
                for row in cfg["chunks"]:
                    print(f"    per chunk k={row['k_per_passage']} cap={row['cap']}: local "
                          f"{row['local_recall']} doc {row['doc_recall']} "
                          f"({row['candidates_per_chunk']} candidates x {row['chunks_per_doc']} chunks)")
                print(f"{corpus_kind:<11} {method:<7} gate={cfg['keyword_gate']:<3} "
                      f"kept {cfg['coverage']['mentions_in_kept_passage']}  "
                      + "  ".join(f"k={r['k']}: P {r['parent']['passage_recall']}"
                                  f" W {r['parent']['window_recall']} D {r['parent']['doc_recall']}"
                                  f" ({r['parent']['unique_ids_per_doc']} ids)" for r in rows),
                      flush=True)
    name = f"retrieval-{args.split}" + (f"-{args.name}" if args.name else "")
    _write(OUT / f"{name}.json", {
        "split": args.split, "documents": len(docs), "excluded_cited_reports": len(excl),
        "attack_mapping": cat.summary(),
        # The corpus follows ATTACK_RELEASE, the mapping --attack-release.
        "corpus_release": attack_data.current().version,
        "unit": "passage = Stage 2c sentence; k = distinct technique ids per passage",
        "results": results})
    print(f"written to {OUT / (name + '.json')}")


# ── support-sample ───────────────────────────────────────────────────────────

def cmd_support_sample(args) -> None:
    s = per_doc_scores(args.name)
    rows = []
    for doc_id in s["rows"]:
        d = s["docs"][doc_id]
        spans: dict[str, list[tuple[int, int]]] = {}
        for m in d.techniques:
            c = s["canonical"](m.label)
            if c and m.start is not None and m.end is not None:
                spans.setdefault(c.split(".", 1)[0], []).append((m.start, m.end))
        rows += evidence.check(doc_id, d.text, s["preds"][doc_id]["ttps"], spans, s["canonical"])
    path = OUT / "runs" / args.name / "support-sample.csv"
    n = evidence.support_sample(rows, {k: v.text for k, v in s["docs"].items()}, path, n=args.n)
    print(f"{n} quotes to judge -> {path}")


# ── in-house set ─────────────────────────────────────────────────────────────

def _run_file(path: Path, disabled: set[str]):
    from pipeline.orchestrator import Document, RunOptions, run_document

    # Stage 4 stays on: relationships are scored as the bundle ships them,
    # graph completion (4b) and long-distance inference (4c) included.
    options = RunOptions.from_env(disabled={"5"} | disabled, required={"3"})
    return run_document(Document(file_path=str(path), original_filename=path.name), options)


def cmd_preannotate(args) -> None:
    import shutil

    src = Path(args.file)
    folder = inhouse.ROOT / args.doc_id
    gold = folder / "gold.json"
    if gold.exists():
        sys.exit(f"{gold} exists — not overwriting an annotation")
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"source{src.suffix.lower()}"
    shutil.copyfile(src, target)
    result = _run_file(target, set())
    _write(gold, inhouse.draft_from_run(args.doc_id, target.name, _predictions(result)))
    print(f"draft (status: pre-annotation) -> {gold}\n"
          "Correct it with docs/eval/annotation-guide.md, fill `layers`, then set status to gold.")


def cmd_inhouse_run(args) -> None:
    from api.run_config import build_manifest

    disabled = {s.strip() for s in (args.disable or "").split(",") if s.strip()}
    run_dir = OUT / "inhouse-runs" / args.name
    _stay_on_the_release(run_dir / "run.json", args.name)
    if not (run_dir / "run.json").exists():
        _write(run_dir / "run.json", {"name": args.name, "disabled": sorted(disabled),
                                      "manifest": build_manifest()})
    for d in inhouse.load():
        out = run_dir / "docs" / f"{d.doc_id}.json"
        if out.exists():
            continue
        t0 = time.monotonic()
        try:
            result = _run_file(d.source, disabled)
            record = {"doc_id": d.doc_id, "seconds": round(time.monotonic() - t0, 1),
                      "stages": result.stage_report(), **_predictions(result)}
        except attack_data.UnknownAttackRelease:
            raise   # the setting is wrong for the whole run, not for this document
        except Exception as exc:
            record = {"doc_id": d.doc_id, "error": f"{type(exc).__name__}: {exc}"}
        _write(out, record)
        print(f"{d.doc_id}: {'ERROR' if 'error' in record else 'ok'}", flush=True)


def cmd_inhouse_score(args) -> None:
    run_dir = OUT / "inhouse-runs" / args.name
    meta = json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    canonical, _ = _canonical_fn(run_release(meta))
    docs = {d.doc_id: d for d in inhouse.load()}
    totals: dict[str, metrics.Counts] = {}
    negatives = [0, 0]
    per_doc = {}
    for p in sorted((run_dir / "docs").glob("*.json")):
        pred = json.loads(p.read_text(encoding="utf-8"))
        d = docs.get(p.stem)
        if d is None or "error" in pred:
            continue
        s = inhouse.score(d, pred, canonical)
        per_doc[d.doc_id] = s
        for layer in ("iocs",):
            if layer in s:
                totals[layer] = totals.get(layer, metrics.Counts()) + s[layer]
        if "techniques" in s:
            totals["techniques"] = totals.get("techniques", metrics.Counts()) + s["techniques"].technique
        if "relations" in s:
            totals["relations"] = totals.get("relations", metrics.Counts()) + s["relations"]["strict"]
        for et, c in s.get("entities", {}).items():
            totals[f"entity:{et}"] = totals.get(f"entity:{et}", metrics.Counts()) + c
        negatives[0] += s["negatives"]["cases"]
        negatives[1] += s["negatives"]["violated"]
    report = {"run": args.name, "documents": len(per_doc),
              "micro": {k: v.as_dict() for k, v in sorted(totals.items())},
              "negatives": {"cases": negatives[0], "violated": negatives[1]}}
    _write(run_dir / "score.json", report)
    print(json.dumps(report, indent=2))


def cmd_hard_docs(args) -> None:
    """Features that make a report hard to ingest, per file: pages scanned,
    tables, images, a second text column.  To pick the Phase 5 documents."""
    import pdfplumber

    from pipeline.stage1_ingestion import _page_is_scanned

    rows = []
    for folder in args.folders:
        for f in sorted(Path(folder).iterdir()):
            row: dict = {"file": str(f), "type": f.suffix.lower()}
            try:
                if f.suffix.lower() == ".pdf":
                    with pdfplumber.open(f) as pdf:
                        pages = pdf.pages
                        row["pages"] = len(pages)
                        row["scanned_pages"] = sum(1 for pg in pages if _page_is_scanned(pg))
                        row["tables"] = sum(len(pg.find_tables()) for pg in pages)
                        row["images"] = sum(len(pg.images) for pg in pages)
                        row["two_column_pages"] = sum(1 for pg in pages if _two_columns(pg))
                elif f.suffix.lower() == ".docx":
                    import docx
                    d = docx.Document(str(f))
                    row["tables"] = len(d.tables)
                    row["images"] = len(d.inline_shapes)
                else:
                    continue
            except Exception as exc:
                row["error"] = str(exc)
            row["hardness"] = (row.get("scanned_pages", 0) * 3 + row.get("tables", 0) * 2
                               + row.get("two_column_pages", 0) + min(row.get("images", 0), 10))
            rows.append(row)
    rows.sort(key=lambda r: -r["hardness"])
    _write(OUT / "hard-docs.json", rows)
    for r in rows:
        print(f"{r['hardness']:>4}  {r['file']}  " + ", ".join(
            f"{k}={v}" for k, v in r.items() if k not in ("file", "hardness", "type")))


def cmd_agreement(args) -> None:
    a = json.loads(Path(args.a).read_text(encoding="utf-8"))
    b = json.loads(Path(args.b).read_text(encoding="utf-8"))
    result = inhouse.agreement(a, b)
    if not result:
        sys.exit("no layer annotated in both files")
    for layer, c in result.items():
        print(f"  {layer:<11} F1 {c['f1']:.3f}  (both {c['tp']}, only b {c['fp']}, only a {c['fn']})")


def _two_columns(page) -> bool:
    """Most text lines start in two distinct horizontal bands."""
    words = page.extract_words()
    if len(words) < 80:
        return False
    mid = float(page.width) / 2
    left = sum(1 for w in words if float(w["x1"]) < mid - 10)
    right = sum(1 for w in words if float(w["x0"]) > mid + 10)
    return min(left, right) > 0.3 * len(words)


def _attack_release_option(p, default: str) -> None:
    p.add_argument("--attack-release", help="ATT&CK release the gold ids are mapped with, e.g. 20.0 "
                   f"(default: {default})")


def main() -> None:
    # Before the first ATT&CK file is read: ATTACK_RELEASE may be set in .env,
    # and a variable already set in the environment still wins (as in main.py).
    load_dotenv(_ROOT / ".env")
    ap = argparse.ArgumentParser(prog="python -m evaluation", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("prepare")
    p.add_argument("--corpus", action="append", help="extra folder of .txt reports to check for copies")
    p.add_argument("--threshold", type=float, default=0.5)
    _attack_release_option(p, "the default track; another one is written to attack-mapping-<release>.json")
    p.set_defaults(fn=cmd_prepare)

    p = sub.add_parser("run")
    p.add_argument("--split", choices=annoctr.SPLITS, default="dev")
    p.add_argument("--name", required=True, help="run name, e.g. baseline-dev")
    p.add_argument("--disable", help="extra stage ids to turn off, e.g. 2d,2e")
    p.add_argument("--allow-degraded", action="store_true", help="score even if the LLM cannot run")
    p.add_argument("--limit", type=int)
    p.set_defaults(fn=cmd_run)

    p = sub.add_parser("score")
    p.add_argument("--name", required=True)
    _attack_release_option(p, "the run's; another one is written to score-attack-<release>.json")
    p.set_defaults(fn=cmd_score)

    p = sub.add_parser("ttp-errors")
    p.add_argument("--name", required=True)
    _attack_release_option(p, "the run's; another one is written to ttp-errors-attack-<release>.json")
    p.set_defaults(fn=cmd_ttp_errors)

    p = sub.add_parser("compare")
    p.add_argument("--a", required=True, help="reference run")
    p.add_argument("--b", required=True, help="candidate run")
    p.add_argument("--task", default="ttp", help="ttp, or an entity type (malware, threat_actor, ...)")
    p.add_argument("--resamples", type=int, default=2000)
    _attack_release_option(p, "the release both runs ran on; required when they differ")
    p.set_defaults(fn=cmd_compare)

    p = sub.add_parser("attack-diff")
    p.add_argument("--old", required=True, help="release of data/attack-stix/v<old>/, e.g. 19.2")
    p.add_argument("--new", required=True, help="release of data/attack-stix/v<new>/, e.g. 20.0")
    p.set_defaults(fn=cmd_attack_diff)

    p = sub.add_parser("release-effect")
    p.add_argument("--a", required=True, help="reference run")
    p.add_argument("--b", required=True, help="candidate run")
    p.add_argument("--old", help="the older release (data/attack-stix/v<old>/); default: the one --a ran on, "
                   "required when both runs ran on the same")
    p.add_argument("--new", help="the newer release (data/attack-stix/v<new>/); default: the one --b ran on, "
                   "required when both runs ran on the same")
    _attack_release_option(p, "the release both runs ran on; required when they differ")
    p.set_defaults(fn=cmd_release_effect)

    p = sub.add_parser("retrieval")
    p.add_argument("--split", choices=annoctr.SPLITS, default="dev")
    p.add_argument("--k", default="1,5,10,20,25")
    p.add_argument("--corpus", default="short",
                   help="comma list of short, description, procedure, both (ADR-0072)")
    p.add_argument("--method", default="dense", help="comma list of dense, bm25, minrank, rrf")
    p.add_argument("--gate", default="on", help="keyword gate: on, off or on,off")
    p.add_argument("--corpus-dir", help="a corpus built with --retrieval-out (default pipeline/data)")
    p.add_argument("--keep-cited", action="store_true",
                   help="keep the procedures MITRE wrote from AnnoCTR reports")
    p.add_argument("--chunk", help="also measure per-chunk candidate lists, "
                   "comma list of k_per_passage:cap, e.g. 5:25,10:25")
    p.add_argument("--name", help="suffix of the output file")
    p.add_argument("--limit", type=int)
    _attack_release_option(p, "the release ATTACK_RELEASE selects")
    p.set_defaults(fn=cmd_retrieval)

    p = sub.add_parser("support-sample")
    p.add_argument("--name", required=True)
    p.add_argument("--n", type=int, default=50)
    p.set_defaults(fn=cmd_support_sample)

    p = sub.add_parser("preannotate")
    p.add_argument("file")
    p.add_argument("--doc-id", required=True)
    p.set_defaults(fn=cmd_preannotate)

    p = sub.add_parser("inhouse-run")
    p.add_argument("--name", required=True)
    p.add_argument("--disable")
    p.set_defaults(fn=cmd_inhouse_run)

    p = sub.add_parser("inhouse-score")
    p.add_argument("--name", required=True)
    p.set_defaults(fn=cmd_inhouse_score)

    p = sub.add_parser("agreement")
    p.add_argument("a")
    p.add_argument("b")
    p.set_defaults(fn=cmd_agreement)

    p = sub.add_parser("hard-docs")
    p.add_argument("folders", nargs="+")
    p.set_defaults(fn=cmd_hard_docs)

    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
